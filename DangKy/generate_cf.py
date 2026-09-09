import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from copy import copy
import pandas as pd
import re

# Load DangKy.xlsx to get email
try:
    df_dangky = pd.read_excel(r'T:\ThuongThuongWEB\DangKy\DangKy.xlsx')
    email_dict = dict(zip(df_dangky['Họ tên học sinh'].str.strip().str.lower(), df_dangky['Địa chỉ email']))
except Exception as e:
    print(f"Error loading DangKy.xlsx: {e}")
    email_dict = {}

# 1. Update the Sample Template
wb = openpyxl.load_workbook(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_Sample.xlsx')
sheet_goc = wb['10T0_Goc']
template = wb['Lê Trung Hiếu']

# Set up Mã số phiếu in G1 and H1
template['G1'] = "Mã số phiếu (Dòng):"
template['G1'].font = Font(bold=True)
template['H1'] = 9 # Row 9 for Lê Trung Hiếu

# Update formulas to use INDEX and $H$1
template['A2'] = "=INDEX('10T0_Goc'!C:C, $H$1)"
template['B5'] = "=INDEX('10T0_Goc'!D:D, $H$1)"
template['B6'] = "=INDEX('10T0_Goc'!A:A, $H$1)"
template['B7'] = "=INDEX('10T0_Goc'!E:E, $H$1)"

item_cols = {
    12: 'F', # Kaki nam
    13: 'H', # Quần sooc nam
    14: 'J', # Chân váy 2 lớp nữ
    15: 'L', # Sơ mi cộc tay
    16: 'P', # Sơ mi dài tay
    17: 'T', # Áo phông trắng
    18: 'V', # Áo phông xanh đen
    19: 'X'  # Bộ thể thao hè
}

for row_idx, col_letter in item_cols.items():
    # Remove any static strikethrough from previous runs
    if template[f'A{row_idx}'].font:
        font = copy(template[f'A{row_idx}'].font)
        font.strike = False
        template[f'A{row_idx}'].font = font
        
    # Ensure checkboxes exist
    template[f'D{row_idx}'] = "☐"
    template[f'E{row_idx}'] = "☐"
    template[f'F{row_idx}'] = "☐"
    
    # Qty & Size with INDEX
    template[f'B{row_idx}'] = f'=IF(ISBLANK(INDEX(\'10T0_Goc\'!{col_letter}:{col_letter}, $H$1)), "", INDEX(\'10T0_Goc\'!{col_letter}:{col_letter}, $H$1))'
    size_col = chr(ord(col_letter) + 1)
    template[f'C{row_idx}'] = f'=IF(ISBLANK(INDEX(\'10T0_Goc\'!{size_col}:{size_col}, $H$1)), "", INDEX(\'10T0_Goc\'!{size_col}:{size_col}, $H$1))'

# Add Conditional Formatting
strikethrough_font = Font(strike=True)
rule1 = FormulaRule(formula=['$B12=""'], stopIfTrue=True, font=strikethrough_font)
template.conditional_formatting.add('A12:A19', rule1)

white_font = Font(color="FFFFFF")
rule2 = FormulaRule(formula=['$B12=""'], stopIfTrue=True, font=white_font)
template.conditional_formatting.add('D12:F19', rule2)

wb.save(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_Sample.xlsx')

# 2. Generate final file
wb_hc = openpyxl.load_workbook(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_Sample.xlsx')
sheet_goc_hc = wb_hc['10T0_Goc']
template_hc = wb_hc['Lê Trung Hiếu']
max_row = sheet_goc_hc.max_row

# Remove old generated sheets if we are appending, but better to clear existing student sheets
# Or just keep it simple, since we start from Sample, it only has Goc and Lê Trung Hiếu

for target_row in range(9, max_row + 1):
    student_name = sheet_goc_hc[f'C{target_row}'].value
    if not student_name or str(student_name).strip() == "":
        continue
        
    if target_row == 9:
        ws = template_hc
        email = email_dict.get(str(student_name).strip().lower(), "")
        ws['B4'] = email
    else:
        sheet_title = str(student_name).strip()
        sheet_title = re.sub(r'[\\*?:/\[\]]', '', sheet_title)[:31]
        
        original_title = sheet_title
        counter = 1
        while sheet_title in wb_hc.sheetnames:
            sheet_title = f"{original_title[:28]}_{counter}"
            counter += 1
            
        ws = wb_hc.copy_worksheet(template_hc)
        ws.title = sheet_title
        
        # Magic: Just update H1 and Email!
        ws['H1'] = target_row
        email = email_dict.get(str(student_name).strip().lower(), "")
        ws['B4'] = email

wb_hc.save(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_HoanChinh.xlsx')
print("Completed dynamic CF generation.")
