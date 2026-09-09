import openpyxl
import io
import sys

with io.open('inspect_a2.txt', 'w', encoding='utf-8') as f:
    wb = openpyxl.load_workbook(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_Sample.xlsx')
    ws = wb['Lê Trung Hiếu']
    f.write(f"A2: {ws['A2'].value}\n")
    f.write(f"A1: {ws['A1'].value}\n")
    f.write(f"A3: {ws['A3'].value}\n")
