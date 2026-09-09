import openpyxl
import pandas as pd
import re

# Load DangKy.xlsx to get email
try:
    df_dangky = pd.read_excel(r'T:\ThuongThuongWEB\DangKy\DangKy.xlsx')
    email_dict = dict(zip(df_dangky['Họ tên học sinh'].str.strip().str.lower(), df_dangky['Địa chỉ email']))
except Exception as e:
    email_dict = {}

# 1. Update the Sample Template
wb = openpyxl.load_workbook(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_Sample.xlsm', keep_vba=True)
sheet_goc = wb['10T0_Goc']
template = wb[wb.sheetnames[1]]

# Fix row 15 (Sơ mi cộc tay)
template['B15'] = '=IF(INDEX(\'10T0_Goc\'!E:E, $H$1)="Nữ", IF(ISBLANK(INDEX(\'10T0_Goc\'!N:N, $H$1)), "", INDEX(\'10T0_Goc\'!N:N, $H$1)), IF(ISBLANK(INDEX(\'10T0_Goc\'!L:L, $H$1)), "", INDEX(\'10T0_Goc\'!L:L, $H$1)))'
template['C15'] = '=IF(INDEX(\'10T0_Goc\'!E:E, $H$1)="Nữ", IF(ISBLANK(INDEX(\'10T0_Goc\'!O:O, $H$1)), "", INDEX(\'10T0_Goc\'!O:O, $H$1)), IF(ISBLANK(INDEX(\'10T0_Goc\'!M:M, $H$1)), "", INDEX(\'10T0_Goc\'!M:M, $H$1)))'

# Fix row 16 (Sơ mi dài tay)
template['B16'] = '=IF(INDEX(\'10T0_Goc\'!E:E, $H$1)="Nữ", IF(ISBLANK(INDEX(\'10T0_Goc\'!R:R, $H$1)), "", INDEX(\'10T0_Goc\'!R:R, $H$1)), IF(ISBLANK(INDEX(\'10T0_Goc\'!P:P, $H$1)), "", INDEX(\'10T0_Goc\'!P:P, $H$1)))'
template['C16'] = '=IF(INDEX(\'10T0_Goc\'!E:E, $H$1)="Nữ", IF(ISBLANK(INDEX(\'10T0_Goc\'!S:S, $H$1)), "", INDEX(\'10T0_Goc\'!S:S, $H$1)), IF(ISBLANK(INDEX(\'10T0_Goc\'!Q:Q, $H$1)), "", INDEX(\'10T0_Goc\'!Q:Q, $H$1)))'

wb.save(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_Sample.xlsm')

# 2. Re-Generate final file HoanChinh
wb_hc = openpyxl.load_workbook(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_Sample.xlsm', keep_vba=True)
sheet_goc_hc = wb_hc['10T0_Goc']
template_hc = wb_hc[wb_hc.sheetnames[1]]
template_row = template_hc['H1'].value
max_row = sheet_goc_hc.max_row

for target_row in range(9, max_row + 1):
    student_name = sheet_goc_hc[f'C{target_row}'].value
    if not student_name or str(student_name).strip() == "":
        continue
        
    if target_row == template_row:
        ws = template_hc
        email = email_dict.get(str(student_name).strip().lower(), "")
        ws['B4'] = email
    else:
        sheet_title = str(student_name).strip()
        sheet_title = re.sub(r'[\\*?:/\[\]]', '', sheet_title)[:31]
        
        original_title = sheet_title
        counter = 1
        while sheet_title in wb_hc.sheetnames:
            sheet_title = f"{original_title[:28]}_{counter}"
            counter += 1
            
        ws = wb_hc.copy_worksheet(template_hc)
        ws.title = sheet_title
        
        ws['H1'] = target_row
        email = email_dict.get(str(student_name).strip().lower(), "")
        ws['B4'] = email

# Save as .xlsm so the macro is preserved in the full file as well!
wb_hc.save(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_HoanChinh.xlsm')
print("Somi formulas fixed and HoanChinh regenerated.")
