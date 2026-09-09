import openpyxl
import io
import sys

with io.open('inspect_out.txt', 'w', encoding='utf-8') as f:
    wb = openpyxl.load_workbook(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_HoanChinh.xlsx')
    f.write(f"Sheets: {wb.sheetnames[:5]}\n")
    ws = wb[wb.sheetnames[1]] # First student sheet
    f.write("Values in Col A:\n")
    for i in range(1, 22):
        f.write(f"A{i}: {ws[f'A{i}'].value}\n")
    f.write("Values in Col B:\n")
    for i in range(1, 10):
        f.write(f"B{i}: {ws[f'B{i}'].value}\n")
