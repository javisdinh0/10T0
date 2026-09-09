import openpyxl
from copy import copy
import pandas as pd
import re

# Load DangKy.xlsx to get email
df_dangky = pd.read_excel(r'T:\ThuongThuongWEB\DangKy\DangKy.xlsx')
email_dict = dict(zip(df_dangky['Họ tên học sinh'].str.strip().str.lower(), df_dangky['Địa chỉ email']))

# Load the user's formatted template
wb = openpyxl.load_workbook(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_Sample.xlsx')
sheet_goc = wb['10T0_Goc']
template_sheet = wb['Lê Trung Hiếu']

item_cols = {
    12: 'F', # Kaki nam
    13: 'H', # Quần sooc nam
    14: 'J', # Chân váy 2 lớp nữ
    15: 'L', # Sơ mi cộc tay
    16: 'P', # Sơ mi dài tay
    17: 'T', # Áo phông trắng
    18: 'V', # Áo phông xanh đen
    19: 'X'  # Bộ thể thao hè
}

# Iterate all students
max_row = sheet_goc.max_row
for target_row in range(9, max_row + 1):
    student_name = sheet_goc[f'C{target_row}'].value
    if not student_name or str(student_name).strip() == "":
        continue
        
    if target_row == 9:
        ws = template_sheet
    else:
        # Create sheet name
        sheet_title = str(student_name).strip()
        sheet_title = re.sub(r'[\\*?:/\[\]]', '', sheet_title)[:31]
        
        original_title = sheet_title
        counter = 1
        while sheet_title in wb.sheetnames:
            sheet_title = f"{original_title[:28]}_{counter}"
            counter += 1
            
        ws = wb.copy_worksheet(template_sheet)
        ws.title = sheet_title

        # FIX: The user moved the name to A2
        ws['A2'] = f"='10T0_Goc'!C{target_row}"
        # Make sure B3 is clear (it was used in my old layout, but here it should be empty)
        ws['B3'] = None
        
        # Update other Formulas
        ws['B5'] = f"='10T0_Goc'!D{target_row}"
        ws['B6'] = f"='10T0_Goc'!A{target_row}"
        ws['B7'] = f"='10T0_Goc'!E{target_row}"
        
        # Update Email
        email = email_dict.get(str(student_name).strip().lower(), "")
        ws['B4'] = email
        
        # Update Items formulas
        for row_idx, col_letter in item_cols.items():
            # Qty formula
            ws[f'B{row_idx}'] = f'=IF(ISBLANK(\'10T0_Goc\'!{col_letter}{target_row}), "", \'10T0_Goc\'!{col_letter}{target_row})'
            # Size formula (next column in goc)
            size_col = chr(ord(col_letter) + 1)
            ws[f'C{row_idx}'] = f'=IF(ISBLANK(\'10T0_Goc\'!{size_col}{target_row}), "", \'10T0_Goc\'!{size_col}{target_row})'

    # Handle Strikethrough and checkboxes based on registration
    for row_idx, col_letter in item_cols.items():
        qty_val = sheet_goc[f"{col_letter}{target_row}"].value
        if qty_val is None or str(qty_val).strip() == "" or str(qty_val).strip() == "0":
            # Not registered
            current_font = ws[f'A{row_idx}'].font
            if current_font:
                new_font = copy(current_font)
                new_font.strike = True
                ws[f'A{row_idx}'].font = new_font
            
            # Clear checkboxes
            ws[f'D{row_idx}'].value = ""
            ws[f'E{row_idx}'].value = ""
            ws[f'F{row_idx}'].value = ""
        else:
            # Registered: ensure no strikethrough (in case template was modified)
            current_font = ws[f'A{row_idx}'].font
            if current_font and current_font.strike:
                new_font = copy(current_font)
                new_font.strike = False
                ws[f'A{row_idx}'].font = new_font
            
            # Ensure checkboxes exist
            ws[f'D{row_idx}'].value = "☐"
            ws[f'E{row_idx}'].value = "☐"
            ws[f'F{row_idx}'].value = "☐"

# Save the final file
wb.save(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_HoanChinh.xlsx')
print("Completed successfully based on sample with correct A2 linking.")
