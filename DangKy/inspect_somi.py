import openpyxl
import io

with io.open('inspect_somi.txt', 'w', encoding='utf-8') as f:
    wb = openpyxl.load_workbook(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_Sample.xlsx')
    ws = wb['10T0_Goc']
    
    # Print row 1 and 2 for columns L to S
    f.write("Columns L to S (index 12 to 19 in openpyxl):\n")
    for col_idx in range(12, 20):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        val1 = ws.cell(row=1, column=col_idx).value
        val2 = ws.cell(row=2, column=col_idx).value
        f.write(f"Col {col_letter}: '{val1}' | '{val2}'\n")
