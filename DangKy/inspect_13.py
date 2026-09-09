import openpyxl
import io

with io.open('inspect_13_qr.txt', 'w', encoding='utf-8') as f:
    wb = openpyxl.load_workbook(r'T:\ThuongThuongWEB\DangKy\10E.xlsx')
    ws = wb.active
    
    # Let's dump Row 13 for context, specifically columns P, Q, R, S
    f.write("Row 13 context:\n")
    f.write(f"C13 (Tên): {ws['C13'].value}\n")
    f.write(f"E13 (Giới tính): {ws['E13'].value}\n")
    f.write(f"P13 (Sơ mi dài tay Nam Qty): {ws['P13'].value}\n")
    f.write(f"Q13 (Sơ mi dài tay Nam Size): {ws['Q13'].value}\n")
    f.write(f"R13 (Sơ mi dài tay Nữ Qty): {ws['R13'].value}\n")
    f.write(f"S13 (Sơ mi dài tay Nữ Size): {ws['S13'].value}\n")
