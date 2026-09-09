import openpyxl
from openpyxl.styles import Font, Alignment
import pandas as pd

# Load DangKy.xlsx to get email
df_dangky = pd.read_excel(r'T:\ThuongThuongWEB\DangKy\DangKy.xlsx')
# Create a dictionary for quick lookup by name (case insensitive)
email_dict = dict(zip(df_dangky['Họ tên học sinh'].str.strip().str.lower(), df_dangky['Địa chỉ email']))

# Load 10T0.xlsx
wb = openpyxl.load_workbook(r'T:\ThuongThuongWEB\DangKy\10T0.xlsx', data_only=False)
sheet_goc = wb.active
sheet_goc.title = '10T0_Goc' # Rename for clear reference

# Row 9 is Lê Trung Hiếu (based on our previous data inspection)
target_row = 9
student_name = sheet_goc[f'C{target_row}'].value

ws = wb.create_sheet(title=student_name)

# Define some styles
font_title = Font(bold=True, color="4169E1", size=14)
font_bold = Font(bold=True)

# 1. Thông tin học sinh
ws['A1'] = "1. Thông tin học sinh"
ws['A1'].font = font_title

ws['A3'] = "Họ và tên:"
ws['A3'].font = font_bold
ws['B3'] = f"='10T0_Goc'!C{target_row}"

ws['A4'] = "Email:"
ws['A4'].font = font_bold
# Lookup email
email = email_dict.get(str(student_name).strip().lower(), "")
ws['B4'] = email

ws['A5'] = "Ngày sinh:"
ws['A5'].font = font_bold
ws['B5'] = f"='10T0_Goc'!D{target_row}"
ws['B5'].number_format = 'DD/MM/YYYY'

ws['A6'] = "Lớp:"
ws['A6'].font = font_bold
ws['B6'] = f"='10T0_Goc'!A{target_row}"

ws['A7'] = "Giới tính:"
ws['A7'].font = font_bold
ws['B7'] = f"='10T0_Goc'!E{target_row}"

# 2. Chi tiết đăng ký
ws['A9'] = "2. Chi tiết đăng ký"
ws['A9'].font = font_title

# Columns in original sheet:
items = [
    ("Kaki nam", "F", "G"),
    ("Quần sooc nam", "H", "I"),
    ("Chân váy 2 lớp nữ", "J", "K"),
    ("Sơ mi cộc tay", "L", "M"),
    ("Sơ mi dài tay", "P", "Q"),
    ("Áo phông trắng", "T", "U"),
    ("Áo phông xanh đen", "V", "W"),
    ("Bộ thể thao hè", "X", "Y")
]

row = 11
# Headers for details
headers = ["Tên món", "Số lượng", "Size", "1", "2", "3"]
for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=row, column=col_idx, value=header)
    cell.font = font_bold
    if col_idx >= 4:
        cell.alignment = Alignment(horizontal='center')

row += 1

# List items using formulas
for item_name, qty_col, size_col in items:
    # Item name
    ws.cell(row=row, column=1, value=item_name)
    # Link to quantity and size using IF to hide blanks
    ws.cell(row=row, column=2, value=f'=IF(ISBLANK(\'10T0_Goc\'!{qty_col}{target_row}), "", \'10T0_Goc\'!{qty_col}{target_row})')
    ws.cell(row=row, column=3, value=f'=IF(ISBLANK(\'10T0_Goc\'!{size_col}{target_row}), "", \'10T0_Goc\'!{size_col}{target_row})')
    
    # Checkboxes
    for col_idx in range(4, 7):
        cell = ws.cell(row=row, column=col_idx, value="☐")
        cell.alignment = Alignment(horizontal='center')
        cell.font = Font(size=14)
    
    row += 1

# Adjust column widths
ws.column_dimensions['A'].width = 25
ws.column_dimensions['B'].width = 30
ws.column_dimensions['C'].width = 15
ws.column_dimensions['D'].width = 8
ws.column_dimensions['E'].width = 8
ws.column_dimensions['F'].width = 8

wb.save(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_Sample.xlsx')
print("Sample created successfully.")
