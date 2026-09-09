import openpyxl
import io
import re

with io.open('scan_errors_10e.txt', 'w', encoding='utf-8') as f:
    wb = openpyxl.load_workbook(r'T:\ThuongThuongWEB\DangKy\10E.xlsx')
    ws = wb.active
    
    qty_cols = ['F', 'H', 'J', 'L', 'N', 'P', 'R', 'T', 'V', 'X']
    size_cols = ['G', 'I', 'K', 'M', 'O', 'Q', 'S', 'U', 'W', 'Y']
    
    for row in range(9, ws.max_row + 1):
        name = ws[f'C{row}'].value
        if not name or str(name).strip() == "":
            continue
            
        gender = str(ws[f'E{row}'].value).strip()
        
        # Check all qty and size cols
        for i in range(len(qty_cols)):
            q_col = qty_cols[i]
            s_col = size_cols[i]
            
            q_val = ws[f'{q_col}{row}'].value
            s_val = ws[f'{s_col}{row}'].value
            
            if q_val is not None and str(q_val).strip() != "":
                q_str = str(q_val).strip()
                # If quantity has letters, it's probably a size
                if re.search(r'[a-zA-Z]', q_str):
                    f.write(f"Row {row} ({name}): Cột Số lượng ({q_col}) chứa chữ cái '{q_str}'. Có thể bị lệch cột.\n")
                    
            if s_val is not None and str(s_val).strip() != "":
                s_str = str(s_val).strip()
                # If size is just a single digit like 1, 2, 3, it might be a qty
                if re.fullmatch(r'\d+', s_str) and int(s_str) < 10:
                    f.write(f"Row {row} ({name}): Cột Size ({s_col}) chứa số nhỏ '{s_str}'. Có thể bị lệch cột.\n")
                    
        # Check gender mismatch for Sơ mi
        # Nam cols: L,M (cộc) and P,Q (dài)
        # Nữ cols: N,O (cộc) and R,S (dài)
        if gender == "Nữ":
            for col in ['L', 'M', 'P', 'Q']:
                val = ws[f'{col}{row}'].value
                if val is not None and str(val).strip() != "":
                    f.write(f"Row {row} ({name}): Nữ nhưng có dữ liệu ở cột dành cho Nam ({col}) là '{val}'.\n")
        elif gender == "Nam":
            for col in ['N', 'O', 'R', 'S', 'J', 'K']: # J,K is váy nữ
                val = ws[f'{col}{row}'].value
                if val is not None and str(val).strip() != "":
                    f.write(f"Row {row} ({name}): Nam nhưng có dữ liệu ở cột dành cho Nữ/Váy ({col}) là '{val}'.\n")
                    
    f.write("Scan complete.\n")
