import openpyxl
import io

with io.open('inspect_layout.txt', 'w', encoding='utf-8') as f:
    wb = openpyxl.load_workbook(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_Sample.xlsx')
    ws = wb['Lê Trung Hiếu']
    for i in range(1, 25):
        f.write(f"A{i}: {ws[f'A{i}'].value} | B{i}: {ws[f'B{i}'].value}\n")
