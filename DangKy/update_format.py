import openpyxl
from copy import copy

wb = openpyxl.load_workbook(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_HoanChinh.xlsx')
sheet_goc = wb['10T0_Goc']

item_cols = {
    12: 'F', # Kaki nam
    13: 'H', # Quần sooc nam
    14: 'J', # Chân váy 2 lớp nữ
    15: 'L', # Sơ mi cộc tay
    16: 'P', # Sơ mi dài tay
    17: 'T', # Áo phông trắng
    18: 'V', # Áo phông xanh đen
    19: 'X'  # Bộ thể thao hè
}

for sheet_name in wb.sheetnames:
    if sheet_name == '10T0_Goc':
        continue
        
    ws = wb[sheet_name]
    
    # Get target_row from B3
    b3_val = ws['B3'].value
    if b3_val and str(b3_val).startswith("='10T0_Goc'!C"):
        try:
            target_row = int(str(b3_val).split("!C")[1])
        except Exception:
            continue
            
        for row_idx, col_letter in item_cols.items():
            qty_val = sheet_goc[f"{col_letter}{target_row}"].value
            
            # If qty is None or empty string or 0, we consider it "not registered"
            if qty_val is None or str(qty_val).strip() == "" or str(qty_val).strip() == "0":
                # Strikethrough item name
                current_font = ws[f'A{row_idx}'].font
                if current_font:
                    new_font = copy(current_font)
                    new_font.strike = True
                    ws[f'A{row_idx}'].font = new_font
                
                # Remove checkboxes
                ws[f'D{row_idx}'].value = ""
                ws[f'E{row_idx}'].value = ""
                ws[f'F{row_idx}'].value = ""

wb.save(r'T:\ThuongThuongWEB\DangKy\10T0_PhieuGiao_HoanChinh.xlsx')
print("Format updated successfully.")
